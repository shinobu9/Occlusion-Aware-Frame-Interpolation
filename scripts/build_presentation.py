"""Create the defense PDF and PPTX from completed metric artifacts, never placeholders."""
import json
from pathlib import Path
import textwrap
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
OUT = ROOT / "presentation"
OUT.mkdir(exist_ok=True)
NAVY, BLUE, ORANGE, GRAY = "#12243B", "#247BA0", "#E98539", "#5D6D7E"
plt.rcParams.update({"font.family": "DejaVu Sans", "pdf.fonttype": 42, "font.size": 14})
slides = []


def slide(title, subtitle=""):
    fig = plt.figure(figsize=(16, 9), facecolor="#F7F9FC")
    fig.text(.055, .91, title, fontsize=30, weight="bold", color=NAVY)
    if subtitle:
        fig.text(.055, .85, subtitle, fontsize=14, color=GRAY)
    fig.text(.055, .035, "Occlusion-Aware Frame Interpolation  |  Kamenskaya · Vokhmyanin · Furmanyuk", fontsize=10, color=GRAY)
    fig.text(.94, .035, f"{len(slides)+1:02d}", fontsize=10, color=GRAY)
    slides.append(fig)
    return fig


def text(fig, x, y, content, size=18, color=NAVY, width=None, **kwargs):
    if width:
        content = "\n".join(textwrap.fill(p, width) for p in content.split("\n"))
    fig.text(x, y, content, fontsize=size, color=color, va="top", linespacing=1.5, **kwargs)


def table(fig, rows, headers, rect=(.06, .25, .88, .5), widths=None, fontsize=17, highlights=()):
    ax = fig.add_axes(rect)
    ax.axis("off")
    tab = ax.table(cellText=rows, colLabels=headers, loc="center", cellLoc="left", colLoc="left", colWidths=widths)
    tab.auto_set_font_size(False)
    tab.set_fontsize(fontsize)
    for (r, c), cell in tab.get_celld().items():
        cell.set_edgecolor("white")
        cell.set_height(1/(len(rows)+1))
        cell.PAD = .08
        if r == 0:
            cell.set_facecolor(NAVY); cell.get_text().set_color("white"); cell.get_text().set_weight("bold")
        else:
            cell.set_facecolor("#DCEEF4" if r-1 in highlights else ("#EAF0F6" if r % 2 else "white"))
            cell.get_text().set_color(NAVY)
    return tab


def main():
    m = pd.read_csv(ART / "mask_summary.csv").set_index("method")
    ci = pd.read_csv(ART / "mask_paired_ci.csv")
    ref = pd.read_csv(ART / "refiner_test_summary.csv").set_index("method")
    diag = pd.read_csv(ART / "refiner_validation_summary.csv")
    runtime = json.loads((ART / "runtime.json").read_text())
    choice = json.loads((ART / "final_selection.json").read_text())
    data = json.loads((ART / "data_report.json").read_text())
    groups = pd.read_csv(ART / "mask_disagreement_subgroups.csv")
    interior_ci = pd.read_csv(ART / "mask_interior_ci.csv")
    interior = interior_ci[(interior_ci.metric == "psnr") & (interior_ci.baseline == "fixed")].iloc[0]
    hname = next(k for k in m.index if k.startswith("heuristic_"))
    delta = m.loc["mask_s42", "psnr"]-m.loc["fixed", "psnr"]
    dci = ci[(ci.method == "mask_s42") & (ci.baseline == "fixed") & (ci.metric == "psnr")].iloc[0]
    oracle_gap = (10**(-m.loc["fixed", "pooled_psnr_dis"]/10)-10**(-m.loc["mask_s42", "pooled_psnr_dis"]/10))/(10**(-m.loc["fixed", "pooled_psnr_dis"]/10)-10**(-m.loc["oracle", "pooled_psnr_dis"]/10))

    fig = slide("Occlusion-Aware Frame Interpolation", "A controlled study of lightweight learned blending")
    text(fig, .06, .72, "Which endpoint should contribute at each pixel?", 28, width=55, weight="bold")
    text(fig, .06, .56, "Two video frames → frozen optical flow → shared warping → blending mask", 22, width=62)
    text(fig, .06, .38, f"{delta:+.3f} dB mean PSNR over fixed blending", 32, BLUE, weight="bold")
    text(fig, .06, .29, "3,782 test triplets · 95k trainable parameters · three mask seeds", 18)
    text(fig, .06, .15, "Lisa Kamenskaya · Pavel Vokhmyanin · Avdey Furmanyuk", 20)

    fig = slide("Related work: established principle, focused experiment", "We claim neither a novel interpolation algorithm nor an exact reproduction of a prior study.")
    table(fig, [
        ["Super SloMo (2018)", "Learned visibility + flow refinement", "We freeze flow; train only the mask"],
        ["Softmax Splatting (2020)", "Weighted forward warping + synthesis", "We use shared weighted splats"],
        ["RIFE (2022)", "Directly learned intermediate flow", "We isolate blending with RAFT-small"],
    ], ["Prior work", "Established idea", "Our narrower scope"], rect=(.06, .34, .88, .4), widths=[.24, .39, .37], fontsize=16)
    text(fig, .06, .24, "Contribution: implementation, controlled comparisons, uncertainty and failure analysis.", 20, BLUE)
    text(fig, .06, .14, "References: Jiang et al., CVPR 2018 · Niklaus & Liu, CVPR 2020 · Huang et al., ECCV 2022\nPaper links and implementation provenance are included in the notebook and README.", 12, GRAY)

    fig = slide("The three methods share the same motion and warps", "Only the rule for choosing the blending mask changes in the primary comparison.")
    text(fig, .08, .74, r"$\hat I_t = M\odot W_0 + (1-M)\odot W_1$", 34, BLUE)
    table(fig, [["Fixed", "M = 0.5 at every pixel", "No training"],
                ["Heuristic", "Flow consistency + splat coverage", "Strength selected on validation"],
                ["Learned", "18-channel input → small U-Net → sigmoid", "95,185 trainable parameters"]],
          ["Method", "Mask", "Learning"], rect=(.06, .32, .88, .31), widths=[.16, .50, .34], fontsize=17, highlights=(2,))
    text(fig, .06, .23, "Frozen RAFT-small · half-flow forward splatting · maps aligned on the midpoint grid", 17)
    text(fig, .06, .14, "A mask can select available information; it cannot reconstruct details missing from both warps.", 17, ORANGE)

    fig = slide("Data integrity before model comparison", "A corrected rerun of exploratory code, with a fixed and auditable protocol.")
    table(fig, [["Training", "1,200", "8k steps; 192×192 crops after full-frame warping"],
                ["Validation", "150", "Checkpoint + heuristic selection; 64 for refiner tuning"],
                ["Test", "3,782", "Full public Vimeo test mirror; never used for selection"]],
          ["Split", "Triplets", "Use"], rect=(.06, .43, .88, .32), widths=[.19, .13, .68], fontsize=17)
    text(fig, .07, .33, "• Pinned dataset revisions, sequence IDs and decoded-frame SHA-256\n• Exact train/test overlap and repeated train/val frames excluded\n• Three mask seeds; identical warps and metrics for every method", 19)
    text(fig, .07, .13, "Directory-prefix separation is not proof of source-video independence. Similar scenes may remain.", 14, GRAY)

    fig = slide("A small learned mask improves the held-out reconstruction", "Mean per-image PSNR/SSIM/MAE; regional PSNR uses pooled squared error.")
    rows = []
    for key, name in [("fixed", "Fixed"), (hname, "Heuristic"), ("mask_s42", "Learned (seed 42)")]:
        r = m.loc[key]
        rows.append([name, f"{r.psnr:.3f}", f"{r.ssim:.5f}", f"{r.mae:.5f}", f"{r.pooled_psnr_dis:.3f}", f"{r.pooled_psnr_occ:.3f}"])
    table(fig, rows, ["Method", "PSNR ↑", "SSIM ↑", "MAE ↓", "PSNR dis ↑", "PSNR occ ↑"], rect=(.06, .43, .88, .31), widths=[.26, .13, .15, .15, .16, .15], fontsize=17, highlights=(2,))
    text(fig, .07, .33, f"ΔPSNR vs fixed: {delta:+.3f} dB  |  paired 95% CI [{dci.ci_lo:+.3f}, {dci.ci_hi:+.3f}]", 23, BLUE, weight="bold")
    gains = m.loc[["mask_s42", "mask_s43", "mask_s44"], "psnr"]-m.loc["fixed", "psnr"]
    text(fig, .07, .25, f"All three mask seeds: gains from {gains.min():+.3f} to {gains.max():+.3f} dB.", 20)
    text(fig, .07, .19, f"Post-hoc check without a 16px border: {interior.delta:+.3f} dB, 95% CI [{interior.ci_lo:+.3f}, {interior.ci_hi:+.3f}].", 17)
    text(fig, .07, .11, "dis / occ are disagreement and occlusion-related proxies, not ground-truth visibility labels.\nBootstrap resamples triplets and does not account for unknown video-level correlations.", 12, GRAY)

    fig = slide("Inspect both gains and failures", "Post-hoc illustrations: 95th-percentile interior gain and 5th-percentile full-frame gain; labels show full-frame ΔPSNR.")
    for rect, name in [((.06, .45, .88, .32), "mask_case_0.png"), ((.06, .10, .88, .32), "mask_case_2.png")]:
        ax = fig.add_axes(rect); ax.imshow(Image.open(ART / name)); ax.axis("off")

    fig = slide("Generative refinement is a separate extension", "Same 3.31M-parameter refiner architecture; residual L1 regression vs conditional flow matching.")
    rows = []
    for key, label in [("learned_cached_reference", "Mask only"), ("l1", "Mask + L1"), ("fm", "Mask + FM"), ("fm_mean4", "FM mean of 4"), ("selected_fm", "Validation-selected FM")]:
        r = ref.loc[key]
        rows.append([label, f"{r.psnr:.3f}", f"{r.ssim:.5f}", f"{r.lpips:.4f}", f"{r.pooled_psnr_dis:.3f}"])
    table(fig, rows, ["Method", "PSNR ↑", "SSIM ↑", "LPIPS ↓", "PSNR dis ↑"], rect=(.06, .30, .88, .46), widths=[.38, .15, .17, .15, .15], fontsize=17)
    conf = choice["config"]
    text(fig, .06, .21, f"Selected on validation: {Path(conf['run']).name} · {conf['solver']} {conf['steps']} · guidance {conf['guidance']:g}", 17, BLUE)
    text(fig, .06, .12, f"FM mean-of-4 improves LPIPS over mask-only but loses PSNR; L1 is better on both.\nSelected FM equals the baseline configuration. Fixed per-triplet noise; paired confidence intervals are in the notebook.", 13, GRAY)

    fig = slide("Diagnose the sampler before increasing training", "Validation only: Euler/Heun × 20/50/100; CFG is evaluated only after condition-dropout training.")
    sam = diag[diag.experiment.str.match(r"val_fm_(euler|heun)\d+$")].sort_values(["solver", "steps"])
    rows = [[r.solver.title(), str(int(r.steps)), f"{r.pooled_psnr_dis:.3f}", f"{r.lpips_window:.4f}", f"{r.refiner_ms_per_triplet:.1f}"] for r in sam.itertuples()]
    table(fig, rows, ["Solver", "Steps", "PSNR dis ↑", "LPIPS window ↓", "ms / triplet*"], rect=(.06, .27, .88, .49), widths=[.20, .13, .22, .25, .20], fontsize=16)
    text(fig, .06, .18, "Additional ablations: EMA → logit-normal time → condition dropout / CFG; 12k vs 30k steps.", 17)
    text(fig, .06, .10, "*Refiner-only batch-4 throughput measurement, not batch-1 end-to-end latency. Heun uses twice as many network calls per step.\nAll ablations and the exact validation selection are preserved in the notebook and CSV files.", 12, GRAY)

    fig = slide("Cost and the limits of a mask-only model", "Runtime is measured on H100 at 256×448, batch size 1, with CUDA synchronization.")
    text(fig, .07, .73, f"{runtime['mask_and_blend_median_ms']:.2f} ms", 39, BLUE, weight="bold")
    text(fig, .07, .62, "mask + blending", 21)
    text(fig, .52, .73, f"{runtime['full_pipeline_median_ms']:.2f} ms", 39, BLUE, weight="bold")
    text(fig, .52, .62, "full two-frame interpolation", 21)
    text(fig, .07, .45, f"The learned mask closes {100*oracle_gap:.1f}% of the oracle MSE gap in disagreement regions.", 23, width=82)
    text(fig, .07, .31, f"• Highest-disagreement quartile: {groups.iloc[-1].delta_psnr:+.3f} dB; gains do not grow monotonically.\n• The oracle bounds convex blending MSE and uses ground truth.\n• Imperfect flow and missing detail limit a mask; FM adds repeated network calls.", 19)
    text(fig, .07, .13, "Reported latency excludes video decoding/encoding. This prototype is not DLSS and makes no gaming real-time claim.", 13, GRAY)

    fig = slide("What this experiment establishes", "A reproducible answer within a deliberately limited setup.")
    text(fig, .07, .73, f"1  Learned blending improves reconstruction: {delta:+.3f} dB vs 50/50.", 24, BLUE, weight="bold")
    text(fig, .07, .59, "2  Gains are consistent across mask seeds, but failures and a large oracle gap remain.", 22, width=85)
    best_label = ref.loc[["l1", "fm", "fm_mean4", "selected_fm"], "psnr"].idxmax()
    text(fig, .07, .43, f"3  Best tested refiner by mean test PSNR: {best_label}.\n    Generative sampling must justify its quality and computational cost.", 22)
    text(fig, .07, .26, "Limitations: small training subset · one refiner seed · proxy visibility regions · unknown scene correlations", 17, GRAY, width=100)
    text(fig, .07, .15, "Deliverables: executed notebook · source + checkpoints + per-triplet CSV · inference demo", 19, BLUE)

    with PdfPages(OUT / "occlusion_aware_frame_interpolation.pdf") as pdf:
        for i, fig in enumerate(slides):
            fig.savefig(OUT / f"slide_{i+1:02d}.png", dpi=140)
            pdf.savefig(fig)
            plt.close(fig)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    for i in range(len(slides)):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(OUT / f"slide_{i+1:02d}.png"), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(OUT / "occlusion_aware_frame_interpolation.pptx")
    thumbs = []
    for i in range(len(slides)):
        im = Image.open(OUT / f"slide_{i+1:02d}.png").convert("RGB")
        im.thumbnail((640, 360)); thumbs.append(im)
    sheet = Image.new("RGB", (1280, 360*5), "white")
    for i, im in enumerate(thumbs): sheet.paste(im, ((i%2)*640, (i//2)*360))
    sheet.save(OUT / "contact_sheet.jpg")
    print(f"Wrote {len(slides)} slides to PDF/PPTX")


if __name__ == "__main__":
    main()
